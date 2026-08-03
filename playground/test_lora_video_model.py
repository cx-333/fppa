# coding: utf-8

# from src.models.image_model import DCVCRTImage
# from src.models.video_model import DCVCRTVideo 
from src.models.lora_image_model import DCVCRTImage
from src.models.lora_video_model import DCVCRTVideo
from src.utils.transforms import ycbcr2rgb, rgb2ycbcr 
from src.utils.utils import get_state_dict, get_padding_size, replicated_pad, AverageMeter

import torch 
from torch import nn 
from torchvision import transforms 
import os 
import argparse 
from piq import ssim, multi_scale_ssim, psnr
import csv 
from PIL import Image 
from tqdm import tqdm 
import openpyxl
import lpips




# load video data: *.png
class PNGVideoDataset(object):
    """
    Dataset/
            video1/... im00001.png
            ...
            videon/... im00001.png
    """
    def __init__(self, root_dir, sep_len: int = 96):
        """
        root_dir: video root directory
        sep_len: frame separation length
        """
        self.root_dir = root_dir 
        self.sep_len = sep_len 
        self.videos = {}
        self.read_videos() 
        
    
    def read_videos(self):
        video_names = [tmp for tmp in os.listdir(self.root_dir) if os.path.isdir(os.path.join(self.root_dir, tmp))]

        for v in video_names:
            self.videos[v] = sorted(os.listdir(os.path.join(self.root_dir, v)), key=lambda x: int(x.lstrip("im").rstrip(".png")))[:self.sep_len]
    

def data_preprocessing(frame_path, half: bool=False):
    img = Image.open(frame_path).convert("RGB")
    img = transforms.ToTensor()(img)
    img = img.unsqueeze(0)
    
    if half:
        img = img.half()
        
    return img



def parse_arg():
    parser = argparse.ArgumentParser(description="Test DCVC-RT Video Model")
    parser.add_argument("--root_dir", type=str, required=True, help="Video root directory")
    parser.add_argument("--dataset", type=str, default="UVG", help="Test dataset name.")
    parser.add_argument("--sep_len", type=int, default=96, help="Frame separation length")
    parser.add_argument("--gop_size", type=int, default=-1, help="Gop size, -1 denotes one Intra frame, remains are P frame.")
    parser.add_argument("--save_dir", type=str, default="./results", help="Save directory")
    parser.add_argument("--save", action="store_true", help="Save results")
    parser.add_argument("--intra_model", type=str, default=None, help="Intra model path")
    parser.add_argument("--intra_lora_model", type=str, default=None, help="Intra lora model path")
    parser.add_argument("--inter_model", type=str, default=None, help="Inter model path")
    parser.add_argument("--inter_lora_model", type=str, default=None, help="Inter lora model path")
    parser.add_argument("--device", type=str, default="cuda", help="Device to use")
    parser.add_argument("--qps", type=int, nargs="+", default=list(range(0, 64, 8)), help="QP value")
    parser.add_argument("--half", action="store_true", help="Use half precision for inference")
    parser.add_argument("--reset_interval", type=int, default=64, help="Reset interval for P feature buffer")
    return parser.parse_args()



@torch.inference_mode()
def main_(args):
    """
    gop_size: gop size, -1 denotes one Intra frame, remains are P frame.
    """
    
    qp_shift = [0, 8, 0, 4, 0, 4, 0, 4]
    
    device = args.device    
    perceptual_loss = lpips.LPIPS(net="alex").to(device)
    
    video_dataset = PNGVideoDataset(args.root_dir, args.sep_len)
    
    image_model = DCVCRTImage().to(device)
    image_model.eval()
    
    video_model = DCVCRTVideo().to(device)
    video_model.eval()
    
    # load model weight
    image_model.load_state_dict(get_state_dict(args.intra_model), strict=False)
    image_model.load_state_dict(get_state_dict(args.intra_lora_model), strict=False)
    video_model.load_state_dict(get_state_dict(args.inter_model), strict=False)
    video_model.load_state_dict(get_state_dict(args.inter_lora_model), strict=False)
    
    if args.half:
        image_model.half()
        video_model.half()

    # collect metrics 
    wb = openpyxl.Workbook()
    wb.remove(wb.active)     # remove default sheet # type: ignore
    qp_avg_metrics = {}   # key: qp, value: {'psnr': avg, 'ssim': avg, 'msssim': avg, 'bpp': avg}
    
    for qp in args.qps:
        
        qp_tr = torch.tensor(qp, dtype=torch.long, device=device)
        
        # init metrics 
        # ---------- 为当前 qp 创建收集列表和 Excel 表 ----------
        sheet_name = f"qp={qp}"
        ws = wb.create_sheet(title=sheet_name)
        # 写表头
        ws.append(["Video Name", "Frame Index", "PSNR", "SSIM", "MS-SSIM", "BPP"])
        
        all_frame_metrics = []   # 临时存放当前 qp 的所有帧指标，用于计算平均值
        
        for video_name, cur_video_frames in tqdm(video_dataset.videos.items(), ncols=120):
            
            print(f"\nTest video {video_name} with qp {qp}.")
            
            # clear reference featur and buffer 
            video_model.clear_dpb()
            # video_model.set_curr_poc(0)
            # print(cur_video_frames)
            
            for i, cur_frame in enumerate(cur_video_frames):
                
                # data preprocessing 
                cur_frame_path = os.path.join(video_dataset.root_dir, video_name, cur_frame)
                
                cur_img = data_preprocessing(cur_frame_path, args.half)
                cur_img = cur_img.to(device)
                
                B, C, H, W = cur_img.size()
                # padding 
                pad_b, pad_r = get_padding_size(H, W, p=64)
                pad_cur_img = replicated_pad(cur_img, pad_b=pad_b, pad_r=pad_r)
                # convert to ycbcr color space 
                pad_cur_img = rgb2ycbcr(pad_cur_img)
                
                # forward 
                if i == 0 or (args.gop_size > 0 and i % args.gop_size == 0):
                    
                    output = image_model.forward(pad_cur_img, qp=qp_tr)

                    video_model.add_ref_frame(feature=None, frame=output["x_hat"])
                
                else:
                    if i % args.reset_interval == 0:
                        video_model.reset_ref_feature() 
                    cur_qp = qp_tr + qp_shift[i % 8]
                    output = video_model.forward(pad_cur_img, qp=cur_qp)
                    
                # convert to rgb color space
                cur_recon = ycbcr2rgb(output["x_hat"])
                
                # compute metrics 
                recon_img = cur_recon[:, :, :H, :W]
                # print(cur_img.min(), cur_img.max(), recon_img.min(), recon_img.max())
                cur_psnr = psnr(cur_img, recon_img, data_range=1.0).item()
                cur_ssim = ssim(cur_img, recon_img, data_range=1.0, full=False).item()      # type: ignore 
                cur_msssim = multi_scale_ssim(cur_img, recon_img, data_range=1.0).item()
                cur_bpp = output["bpp"].mean().item()
                cur_lpips = perceptual_loss(cur_img, recon_img, normalize=True).item()
                
                frame_idx = i + 1 
                # collect metrics 
                all_frame_metrics.append((cur_psnr, cur_ssim, cur_msssim, cur_bpp, cur_lpips))
                ws.append([video_name, frame_idx, cur_psnr, cur_ssim, cur_msssim, cur_bpp])
                
        # ---------- 一个 qp 结束后，计算该 qp 的平均指标并存储 ----------
        if all_frame_metrics:
            avg_psnr = sum(m[0] for m in all_frame_metrics) / len(all_frame_metrics)
            avg_ssim = sum(m[1] for m in all_frame_metrics) / len(all_frame_metrics)
            avg_msssim = sum(m[2] for m in all_frame_metrics) / len(all_frame_metrics)
            avg_bpp = sum(m[3] for m in all_frame_metrics) / len(all_frame_metrics)
            avg_lpips = sum(m[4] for m in all_frame_metrics) / len(all_frame_metrics)
        else:
            avg_psnr = avg_ssim = avg_msssim = avg_bpp = avg_lpips = 0.0
        qp_avg_metrics[qp] = {
            'psnr': avg_psnr,
            'ssim': avg_ssim,
            'msssim': avg_msssim,
            'bpp': avg_bpp,
            'lpips': avg_lpips
        }
        print(f"QP={qp}: Avg PSNR={avg_psnr:.6f}, Avg SSIM={avg_ssim:.6f}, Avg MS-SSIM={avg_msssim:.6f}, Avg BPP={avg_bpp:.6f}, Avg LPIPS: {avg_lpips:.6f}")
        
    # ========== 所有 qp 处理完毕，创建汇总表 ==========
    ws_summary = wb.create_sheet(title="Summary", index=0)   # 放到最前面
    ws_summary.append(["QP", "Avg PSNR", "Avg SSIM", "Avg MS-SSIM", "Avg BPP", "Avg LPIPS"])
    for qp, metrics in qp_avg_metrics.items():
        ws_summary.append([qp, metrics['psnr'], metrics['ssim'], metrics['msssim'], metrics['bpp'], metrics['lpips']])

    # 保存 Excel
    excel_path = os.path.join(args.save_dir, f"{args.dataset}_metrics_results.xlsx")
    wb.save(excel_path)
    print(f"Excel results saved to {excel_path}")

    # ========== 保存 TXT 总体指标 ==========
    txt_path = os.path.join(args.save_dir, f"{args.dataset}_summary.txt")
    with open(txt_path, 'w') as f:
        f.write("QP\tAvg PSNR\tAvg SSIM\tAvg MS-SSIM\tAvg BPP\tAvg LPIPS\n")
        for qp, metrics in qp_avg_metrics.items():
            f.write(f"{qp}\t{metrics['psnr']:.6f}\t{metrics['ssim']:.6f}\t{metrics['msssim']:.6f}\t{metrics['bpp']:.6f}\t{metrics['lpips']:.6f}\n")
    print(f"Summary saved to {txt_path}")    
    print("Test done.")    
    

if __name__ == "__main__":
    args = parse_arg()
    main_(args)
    
    
    